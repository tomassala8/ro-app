#!/usr/bin/env python3
"""
fuentes_sueldos/generar_sueldos.py · sueldos del equipo (E0 ronda 5, NUEVA REGLA DE TOMÁS del 2-oct).

Regla: los sueldos SÍ entran en la app, visibles SOLO para dirección (Tomás) y RRHH (Cecilia). Ni Mili, ni Sofía,
ni nadie más. En «ver como» mandan los permisos de la persona vista (Tomás viendo como Mili no los ve).

Lee (solo lectura): ~/Downloads/SALARIOS EQUIPO (4).xlsx
  · hojas por mes (ENERO 2026 … SEPTIEMBRE 2026): rol, persona, pagado, euros, salario, bonus, condiciones
  · «Proyección Salarial 2026»: salario base, euros, salario por hora, bonus, notas de Tomás y RRHH
  · «Análisis»: inversión por área
  · la hoja «Cuentas bancarias» NO SE ABRE NUNCA (ni se lee ni se copia).
Escribe: data/sueldos/_privado/sueldos.json — almacén privado: solo se abre con /api/ver_dato (regla «sueldos»)
y cada apertura queda en el rastro. No se sirve nunca como fichero ni por /api/modulo. Fuera: correos de contacto.
"""
import json
import re
import sys
import unicodedata
from datetime import datetime
from pathlib import Path

AQUI = Path(__file__).resolve().parent.parent
ORIGEN = Path.home() / "Downloads/SALARIOS EQUIPO (4).xlsx"
SALIDA = AQUI / "data/sueldos/_privado/sueldos.json"
MESES = ["ENERO", "FEBRERO", "MARZO", "ABRIL", "MAYO", "JUNIO", "JULIO", "AGOSTO", "SEPTIEMBRE", "OCTUBRE", "NOVIEMBRE", "DICIEMBRE"]
HOJAS_PROHIBIDAS = re.compile(r"banc", re.I)


def norm(t):
    t = unicodedata.normalize("NFKD", str(t or "")).encode("ascii", "ignore").decode().lower()
    return re.sub(r"[^a-z0-9@.]+", " ", t).strip()


def main():
    import openpyxl
    personas = json.loads((AQUI / "data/personas.json").read_text())
    por_correo = {}
    por_nombre = {}
    for p in personas:
        for c in [p.get("correo"), *(p.get("otros_correos") or [])]:
            if c:
                por_correo[c.strip().lower()] = p["id"]
        por_nombre[norm(p["nombre"])] = p["id"]
        por_nombre.setdefault(norm(p["nombre"].split()[0]), p["id"])
        for a in [p.get("alias"), *(p.get("alias_todos") or [])]:
            if a:
                por_nombre.setdefault(norm(a), p["id"])

    def quien(nombre, correo=None):
        if correo and str(correo).strip().lower() in por_correo:
            return por_correo[str(correo).strip().lower()]
        n = norm(nombre)
        return por_nombre.get(n) or por_nombre.get(n.split()[0] if n else "")

    wb = openpyxl.load_workbook(ORIGEN, data_only=True, read_only=True)
    hojas = [n for n in wb.sheetnames if not HOJAS_PROHIBIDAS.search(n)]   # «Cuentas bancarias» ni se abre
    out = {"personas": {}, "areas": {}, "sin_emparejar": []}

    def fila_dict(cab, fila):
        return {str(c).strip(): v for c, v in zip(cab, fila) if c}

    for nombre_hoja in hojas:
        ws = wb[nombre_hoja]
        filas = [r for r in ws.iter_rows(values_only=True) if any(c is not None for c in r)]
        if not filas:
            continue
        cab = [str(c).strip() if c else None for c in filas[0]]
        mes = next((i + 1 for i, m in enumerate(MESES) if nombre_hoja.upper().startswith(m)), None)
        if mes:
            clave_mes = f"2026-{mes:02d}"
            for r in filas[1:]:
                d = fila_dict(cab, r)
                if not isinstance(d.get("Persona"), str) or not d["Persona"].strip() or d["Persona"].strip().lower().startswith("salario"):
                    continue          # filas vacías o el segundo bloque de la hoja («Salario pactado»): no son personas
                pid = quien(d.get("Persona"), d.get("mail contacto"))
                if not pid:            # persona que no está en personas.json (p. ej. ya fuera): se guarda por su nombre
                    pid = "nombre:" + norm(d["Persona"]).replace(" ", "_")
                    out["sin_emparejar"].append({"hoja": nombre_hoja, "persona": d["Persona"].strip()})
                reg = out["personas"].setdefault(pid, {"meses": {}, "proyeccion": None})
                reg["meses"][clave_mes] = {
                    "rol": (d.get("ROL") or "").strip() or None, "pagado": d.get("Pagado"),
                    "salario_usd": d.get("Salario"), "euros": d.get("EUROS"), "bonus": d.get("BONUS"),
                    "condiciones": d.get("Condiciones especiales"), "detalle": d.get("Detalle"),
                }
        elif nombre_hoja.lower().startswith("proyecci"):
            for r in filas[1:]:
                d = fila_dict(cab, r)
                if not isinstance(d.get("Persona"), str) or not d["Persona"].strip() or d["Persona"].strip().lower() == "persona":
                    continue
                if len(d["Persona"]) > 40:          # notas al pie de la hoja, no personas
                    out.setdefault("notas_proyeccion", []).append(d["Persona"].strip())
                    continue
                pid = quien(d.get("Persona")) or "nombre:" + norm(d["Persona"]).replace(" ", "_")
                out["personas"].setdefault(pid, {"meses": {}, "proyeccion": None})["proyeccion"] = {
                    "rol": d.get("Rol"), "salario_base_usd": d.get("Salario base"), "euros": d.get("Euros"),
                    "salario_hora": d.get("Salario x h (ClickUp)"), "bonus": d.get("Bonus"), "detalle": d.get("Detalle"),
                    "notas": d.get("Notas Tomás / RRHH"),
                }
        elif nombre_hoja.lower().startswith("an"):
            out["areas"]["analisis"] = {"filas": [fila_dict(cab, r) for r in filas[1:] if r and r[1]]}
    out["_meta"] = {"generado": datetime.now().isoformat(timespec="minutes"), "origen": ORIGEN.name,
                    "hojas_leidas": hojas, "hoja_no_abierta": "Cuentas bancarias",
                    "regla": "Solo dirección y RRHH, con /api/ver_dato y rastro (regla «sueldos» de reglas_permisos.json)."}
    SALIDA.parent.mkdir(parents=True, exist_ok=True)
    (SALIDA.parent / ".gitignore").write_text("*\n")
    tmp = SALIDA.with_suffix(".tmp")
    tmp.write_text(json.dumps(out, ensure_ascii=False, indent=1, default=str))
    tmp.replace(SALIDA)
    print(f"sueldos: {len(out['personas'])} personas · {len(hojas)} hojas (sin «Cuentas bancarias») · sin emparejar {len(out['sin_emparejar'])}")
    return 0


if __name__ == "__main__":
    sys.exit(main())
